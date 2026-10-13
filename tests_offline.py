"""Testes offline do SIGI (sem rede):  python tests_offline.py"""
import io
import shutil
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import database as db
import doc_generator as dg
from utils import (
    data_por_extenso, gerar_justificativa, normalizar, nome_sem_tipo,
    preposicao_logradouro, subtrair_meses, TIPOS_SERVICO,
)

tmp = Path(tempfile.mkdtemp())

# --- utils ------------------------------------------------------------------
assert data_por_extenso(date(2026, 6, 2)) == "02 de junho de 2026"
assert subtrair_meses(date(2026, 8, 31), 6) == date(2026, 2, 28)
assert normalizar("Rua Moysés Teixeira-Hausen") == "rua moyses teixeira hausen"
assert preposicao_logradouro("Beco do Sol") == "NO" and preposicao_logradouro("Rua A") == "NA"
assert nome_sem_tipo("Avenida Ville Roy") == "Ville Roy"
for t in TIPOS_SERVICO:
    j = gerar_justificativa(t, "Rua Teste", "Centro")
    assert "Rua Teste" in j and "{" not in j, t
print("utils OK")

# --- template padrão ----------------------------------------------------------
tpl = dg.criar_template_padrao(tmp / "modelo.docx")
rel = dg.analisar_template(tpl)
assert not rel["faltando"], rel
print("template padrão OK:", rel)

# --- substituição com variável QUEBRADA em vários runs + cabeçalho + tabela -------
d = Document()
p = d.add_paragraph()
p.add_run("Local: ")
r1 = p.add_run("{{RU"); r1.bold = True
r2 = p.add_run("A}}"); r2.bold = True
p.add_run(" no bairro ")
p.add_run("{{ BAIRRO }}").italic = True
q = d.add_paragraph(); q.add_run("Justificativa: ").bold = True; q.add_run("{{JUSTIFICATIVA}}")
tab = d.add_table(rows=1, cols=1); tab.rows[0].cells[0].text = "Serviço: {{TIPO_SERVICO}} em {{DATA}}"
d.sections[0].header.paragraphs[0].text = "Cabeçalho {{BAIRRO}}"
quebrado = tmp / "quebrado.docx"; d.save(quebrado)

rel = dg.analisar_template(quebrado)
assert "{{RUA}}" in rel["encontradas"], rel

valores = dg.montar_valores("Rua Moysés Teixeira Hausen", "Caranã", "Tapa-Buraco",
                            "Parágrafo um.\n\nParágrafo dois.", date(2026, 6, 2))
out = Document(io.BytesIO(dg.gerar_documento(quebrado, valores)))
textos = [x.text for x in out.paragraphs]
print(textos)
assert textos[0] == "Local: RUA MOYSÉS TEIXEIRA HAUSEN no bairro CARANÃ", textos[0]
assert out.paragraphs[0].runs[1].bold is True            # formatação preservada
assert textos[1] == "Justificativa: Parágrafo um." and textos[2] == "Parágrafo dois."
assert out.tables[0].rows[0].cells[0].text == "Serviço: TAPA-BURACO em 02 de junho de 2026"
assert out.sections[0].header.paragraphs[0].text == "Cabeçalho CARANÃ"
assert "{{" not in "".join(textos)
print("substituição robusta OK")

# --- lote + zip -----------------------------------------------------------------
itens = [{"rua": f"Rua {i}", "bairro": "Centro", "tipo_servico": "Recapeamento Asfáltico",
          "justificativa": gerar_justificativa("Recapeamento Asfáltico", f"Rua {i}", "Centro")} for i in range(1, 4)]
res = dg.gerar_lote(tpl, itens, tmp / "lote")
assert len(res) == 3 and all(r["caminho"].exists() for r in res)
z = dg.criar_zip([r["caminho"] for r in res])
assert z[:2] == b"PK"
print("lote/zip OK:", [r["arquivo"] for r in res])

# --- modelo sem tags com cabeçalho/rodapé e mapa --------------------------------
from PIL import Image
doc_sem_tags = Document()
sec_st = doc_sem_tags.sections[0]
img_dummy = Image.new("RGB", (60, 60), color="blue")
buf_img = io.BytesIO(); img_dummy.save(buf_img, "PNG"); buf_img.seek(0)
sec_st.header.paragraphs[0].add_run().add_picture(buf_img)
buf_img.seek(0)
sec_st.footer.paragraphs[0].add_run().add_picture(buf_img)
doc_sem_tags.add_paragraph("ESTADO DE RORAIMA - CAMARA MUNICIPAL")
doc_sem_tags.add_paragraph("SOLICITO A EXECUÇÃO DO SERVIÇO DE TAPA-BURACO NA RUA ANTIGA NO BAIRRO CENTRO.")
doc_sem_tags.add_paragraph("Justificativa: Demanda antiga dos moradores.")
doc_sem_tags.add_paragraph("Boa Vista - RR, 01 de janeiro de 2026.")
caminho_st = tmp / "modelo_sem_tags.docx"
doc_sem_tags.save(caminho_st)

rel_st = dg.analisar_template(caminho_st)
assert rel_st["modo"] == "auto", rel_st
assert rel_st["tem_cabecalho_img"] and rel_st["tem_rodape_img"]
assert not rel_st["faltando"], rel_st

buf_mapa = io.BytesIO(); img_dummy.save(buf_mapa, "PNG"); buf_mapa.seek(0)
bytes_doc_gerado = dg.gerar_documento(caminho_st, valores, imagem_mapa=buf_mapa.getvalue(), incluir_mapa=True)
doc_verif = Document(io.BytesIO(bytes_doc_gerado))
textos_verif = " ".join(p.text for p in doc_verif.paragraphs)
assert "RUA MOYSÉS TEIXEIRA HAUSEN" in textos_verif
assert "CARANÃ" in textos_verif
assert sum(1 for r in doc_verif.sections[0].header.part.rels.values() if "image" in r.reltype.lower()) == 1
assert sum(1 for r in doc_verif.sections[0].footer.part.rels.values() if "image" in r.reltype.lower()) == 1
print("modelo oficial sem tags com imagens OK")


# --- banco ---------------------------------------------------------------------
db.DB_PATH = tmp / "t.db"
db.init_db()
db.inserir_indicacao("Rua Moysés Teixeira Hausen", "Caranã", "Tapa-Buraco", date.today() - timedelta(days=30))
db.inserir_indicacao("Rua Velha", "Caranã", "Tapa-Buraco", date.today() - timedelta(days=400))
db.inserir_indicacao("Rua X", "Centro", "Sinalização", date.today())
assert db.buscar_duplicidade("rua moyses teixeira hausen", "Tapa-Buraco", "Caranã")
assert not db.buscar_duplicidade("Rua Moysés Teixeira Hausen", "Sinalização", "Caranã")
assert not db.buscar_duplicidade("Rua Velha", "Tapa-Buraco", "Caranã")
assert db.contar_por_data() == 1
assert ("rua moyses teixeira hausen", "carana") in db.mapa_recentes("Tapa-Buraco")
assert len(db.serie_diaria(14)) == 14
print("banco OK")

shutil.rmtree(tmp, ignore_errors=True)
print("\nTODOS OS TESTES PASSARAM")
